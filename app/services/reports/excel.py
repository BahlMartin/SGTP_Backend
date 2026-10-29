"""
Generador del informe diario asistencial en formato Excel (.xlsx) mediante OpenPyXL.
Responsabilidad única: Compilación en memoria de un libro Excel multi-hoja a partir de las métricas consolidadas.

Estructura de Hojas (Pestañas):
  1. Resumen y Flujo: KPIs generales, SLA de Triage, Franjas Horarias, Curva Horaria y Obras Sociales.
  2. Estudios e Insumos: Resumen por sección bioquímica, logística de tubos y detalle nominal de prácticas.
  3. Personal y Boxes: Rendimiento de Admisión, Técnicos de Box y uso por Box físico (incl. prioridad Box 1).
  4. Horas Extras: Detalle nominal de extensiones horarias y excepciones de turno activas en la jornada.
  5. Auditoría: Indicadores de rectificación en box, estudios agregados/cancelados, cierres ausentes y Triage 'Otro'.
  6. Detalle de Tickets: Sábana completa fila por fila de todos los tickets admitidos en la jornada.
"""
import io
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ==============================================================================
# Paleta institucional y metadatos de Triage (alineados con pdf.py)
# ==============================================================================
INSTITUCION = "Laboratorio de Análisis Clínicos · SGTP"
TITULO_LIBRO = "Informe Diario de Actividad Asistencial"

HEX_TEAL_DARK = "0F766E"
HEX_TEAL_SOFT = "F0FDFA"
HEX_SLATE_DARK = "0F172A"
HEX_SLATE_HEADER = "1E293B"
HEX_SLATE_MUTED = "64748B"
HEX_BG_TH = "E2E8F0"
HEX_BG_ZEBRA = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_RED_SOFT = "FEE2E2"
HEX_RED_TEXT = "B91C1C"
HEX_GREEN_SOFT = "DCFCE7"
HEX_GREEN_TEXT = "15803D"

OBJETIVOS_SLA_TRIAGE: Dict[str, Optional[float]] = {
    'guardia': 5.0,
    'medicos': 10.0,
    'discapacidad': 10.0,
    'oncologia': 15.0,
    'extraccion con turno': 15.0,
    'extraccion sin turno': 30.0,
    'otro': None,
}

ETIQUETAS_TRIAGE: Dict[str, str] = {
    'guardia': 'Guardia / Urgencia Inmediata',
    'medicos': 'Médicos / Personal Sanitario',
    'discapacidad': 'Discapacidad / Movilidad Reducida',
    'oncologia': 'Pacientes Oncológicos',
    'extraccion con turno': 'Extracción con Turno Previo',
    'extraccion sin turno': 'Extracción Espontánea sin Turno',
    'otro': 'Otro Tipo de Atención',
}


def _fecha_legible(fecha_iso: str) -> str:
    if not fecha_iso:
        return 'N/D'
    texto = str(fecha_iso).strip()
    if ' al ' in texto:
        partes = texto.split(' al ', 1)
        return f"{_fecha_legible(partes[0].strip())} - {_fecha_legible(partes[1].strip())}"
    try:
        return datetime.strptime(texto, '%Y-%m-%d').strftime('%d/%m/%Y')
    except (TypeError, ValueError):
        return texto


class ReportExcelGenerator:
    """
    Compila el reporte diario asistencial en un archivo binario Excel (.xlsx) en memoria.
    """

    # Estilos reutilizables
    _FONT_BANNER_TITLE = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    _FONT_BANNER_SUB = Font(name='Calibri', size=10, italic=True, color='E2E8F0')
    _FONT_SECTION = Font(name='Calibri', size=11, bold=True, color=HEX_TEAL_DARK)
    _FONT_TH = Font(name='Calibri', size=10, bold=True, color=HEX_SLATE_DARK)
    _FONT_CELL = Font(name='Calibri', size=10, bold=False, color=HEX_SLATE_DARK)
    _FONT_CELL_BOLD = Font(name='Calibri', size=10, bold=True, color=HEX_SLATE_DARK)
    _FONT_CELL_MUTED = Font(name='Calibri', size=9, italic=True, color=HEX_SLATE_MUTED)
    _FONT_OK = Font(name='Calibri', size=10, bold=True, color=HEX_GREEN_TEXT)
    _FONT_ALERT = Font(name='Calibri', size=10, bold=True, color=HEX_RED_TEXT)

    _FILL_BANNER = PatternFill(start_color=HEX_TEAL_DARK, end_color=HEX_TEAL_DARK, fill_type='solid')
    _FILL_SECTION = PatternFill(start_color=HEX_TEAL_SOFT, end_color=HEX_TEAL_SOFT, fill_type='solid')
    _FILL_TH = PatternFill(start_color=HEX_BG_TH, end_color=HEX_BG_TH, fill_type='solid')
    _FILL_ZEBRA = PatternFill(start_color=HEX_BG_ZEBRA, end_color=HEX_BG_ZEBRA, fill_type='solid')
    _FILL_OK = PatternFill(start_color=HEX_GREEN_SOFT, end_color=HEX_GREEN_SOFT, fill_type='solid')
    _FILL_ALERT = PatternFill(start_color=HEX_RED_SOFT, end_color=HEX_RED_SOFT, fill_type='solid')

    _THIN_SIDE = Side(border_style='thin', color=HEX_BORDER)
    _MED_TEAL_SIDE = Side(border_style='medium', color=HEX_TEAL_DARK)

    _BORDER_CELL = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
    _BORDER_TH = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_MED_TEAL_SIDE)
    _BORDER_SECTION = Border(bottom=_MED_TEAL_SIDE)

    _ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
    _ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')
    _ALIGN_CENTER = Alignment(horizontal='center', vertical='center')

    @classmethod
    def _agregar_encabezado_hoja(
        cls,
        ws: Worksheet,
        subtitulo_hoja: str,
        fecha_txt: str,
        generado_txt: str,
        max_cols: int = 6
    ) -> int:
        """
        Inserta el bloque institucional superior en la hoja activa y retorna la siguiente fila libre.
        """
        ws.views.sheetView[0].showGridLines = True
        max_cols = max(2, max_cols)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_cols)

        c1 = ws.cell(row=1, column=1, value=f"{INSTITUCION} — {TITULO_LIBRO}")
        c2 = ws.cell(
            row=2,
            column=1,
            value=f"{subtitulo_hoja} · Jornada: {fecha_txt} · Generado: {generado_txt} (Hora Argentina)"
        )

        for col_idx in range(1, max_cols + 1):
            cell_top = ws.cell(row=1, column=col_idx)
            cell_sub = ws.cell(row=2, column=col_idx)
            cell_top.fill = cls._FILL_BANNER
            cell_sub.fill = cls._FILL_BANNER

        c1.font = cls._FONT_BANNER_TITLE
        c1.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        c2.font = cls._FONT_BANNER_SUB
        c2.alignment = Alignment(horizontal='left', vertical='center', indent=1)

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 10
        return 4

    @classmethod
    def _agregar_titulo_seccion(cls, ws: Worksheet, fila: int, titulo: str, max_cols: int = 6) -> int:
        """Inserta un subtítulo de sección destacado y retorna la siguiente fila."""
        max_cols = max(2, max_cols)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=max_cols)
        celda = ws.cell(row=fila, column=1, value=titulo)
        celda.font = cls._FONT_SECTION
        celda.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        for c in range(1, max_cols + 1):
            cell_c = ws.cell(row=fila, column=c)
            cell_c.fill = cls._FILL_SECTION
            cell_c.border = cls._BORDER_SECTION
        ws.row_dimensions[fila].height = 22
        return fila + 1

    @classmethod
    def _agregar_tabla(
        cls,
        ws: Worksheet,
        fila_inicio: int,
        encabezados: Sequence[str],
        filas_datos: Sequence[Sequence[Any]],
        alineaciones: Optional[Sequence[str]] = None,
        mensaje_vacio: str = "Sin registros para la jornada seleccionada."
    ) -> int:
        """
        Renderiza una tabla con cabeceras estilizadas, filas cebradas y bordes limpios.
        Retorna la siguiente fila disponible (dejando 1 fila en blanco de separación).
        """
        num_cols = len(encabezados)
        for col_idx, texto_th in enumerate(encabezados, start=1):
            c = ws.cell(row=fila_inicio, column=col_idx, value=texto_th)
            c.font = cls._FONT_TH
            c.fill = cls._FILL_TH
            c.border = cls._BORDER_TH
            align_code = alineaciones[col_idx - 1] if (alineaciones and col_idx - 1 < len(alineaciones)) else 'L'
            c.alignment = (
                cls._ALIGN_RIGHT if align_code == 'R'
                else cls._ALIGN_CENTER if align_code == 'C'
                else cls._ALIGN_LEFT
            )
        ws.row_dimensions[fila_inicio].height = 21

        fila_actual = fila_inicio + 1
        if not filas_datos:
            ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=num_cols)
            c_vacia = ws.cell(row=fila_actual, column=1, value=mensaje_vacio)
            c_vacia.font = cls._FONT_CELL_MUTED
            c_vacia.alignment = cls._ALIGN_LEFT
            for col_idx in range(1, num_cols + 1):
                ws.cell(row=fila_actual, column=col_idx).border = cls._BORDER_CELL
            ws.row_dimensions[fila_actual].height = 19
            return fila_actual + 2

        for idx_fila, valores in enumerate(filas_datos):
            es_par = idx_fila % 2 == 1
            for col_idx, valor in enumerate(valores, start=1):
                celda = ws.cell(row=fila_actual, column=col_idx, value=valor)
                celda.font = cls._FONT_CELL
                celda.border = cls._BORDER_CELL
                if es_par:
                    celda.fill = cls._FILL_ZEBRA

                align_code = alineaciones[col_idx - 1] if (alineaciones and col_idx - 1 < len(alineaciones)) else 'L'
                if align_code == 'R':
                    celda.alignment = cls._ALIGN_RIGHT
                elif align_code == 'C':
                    celda.alignment = cls._ALIGN_CENTER
                else:
                    celda.alignment = cls._ALIGN_LEFT

                if isinstance(valor, float):
                    celda.number_format = '0.0'
                elif isinstance(valor, int):
                    celda.number_format = '#,##0'

            ws.row_dimensions[fila_actual].height = 19
            fila_actual += 1

        return fila_actual + 1

    @classmethod
    def _autoajustar_columnas(cls, ws: Worksheet, min_width: int = 12, max_width: int = 48) -> None:
        """Calcula y asigna un ancho óptimo a cada columna ignorando celdas combinadas de títulos."""
        for col in ws.columns:
            col_idx = col[0].column
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for cell in col:
                if cell.row in (1, 2):
                    continue
                # Evitar medir celdas combinadas horizontales (como subtítulos de sección)
                if any(
                     cell.coordinate in merged_range and merged_range.min_col != merged_range.max_col
                    for merged_range in ws.merged_cells.ranges
                ):
                    continue
                val_str = str(cell.value or '')
                if val_str:
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 4, max_width))

    # ==========================================================================
    # Construcción de cada una de las 6 hojas
    # ==========================================================================
    @classmethod
    def _construir_hoja_resumen_y_flujo(
        cls,
        ws: Worksheet,
        m: Dict[str, Any],
        fecha_txt: str,
        generado_txt: str
    ) -> None:
        ws.title = "1. Resumen y Flujo"
        fila = cls._agregar_encabezado_hoja(
            ws,
            "Hoja 1 · Resumen Ejecutivo, SLA de Triage, Franjas Horarias y Obras Sociales",
            fecha_txt,
            generado_txt,
            max_cols=6
        )

        # 1. KPIs Generales
        fila = cls._agregar_titulo_seccion(ws, fila, "1. Indicadores Clave de la Jornada (KPIs)", max_cols=6)
        filas_kpis = [
            ("Tickets Emitidos (Ingresos Totales)", int(m.get('total_emitidos', 0)), "Total de pacientes admitidos en la fecha"),
            ("Pacientes Atendidos (Finalizados)", int(m.get('total_atendidos', 0)), "Atenciones completadas en box de extracción"),
            ("Pacientes Críticos Atendidos", int(m.get('total_criticos_atendidos', 0)), "Prioridad Guardia + Médicos finalizados"),
            ("Tickets Cancelados / Ausentes", int(m.get('total_cancelados', 0)), "Pacientes ausentes al llamado o dados de baja"),
            ("Tickets En Curso / Pendientes", int(m.get('total_en_curso', 0)), "En sala de espera o en box actualmente"),
            ("Tasa de Completitud Asistencial (%)", float(m.get('tasa_completitud_porcentaje', 0.0)), "Porcentaje de ingresos atendidos con éxito"),
            ("Tasa de Deserción / Cancelación (%)", float(m.get('tasa_desercion_porcentaje', 0.0)), "Porcentaje de ingresos cancelados o ausentes"),
            ("Tiempo Promedio de Espera (min)", float(m.get('promedio_espera_minutos', 0.0)), "Desde admisión hasta llamado en box"),
            ("Tiempo Promedio de Atención en Box (min)", float(m.get('promedio_atencion_minutos', 0.0)), "Desde llamado hasta cierre de atención en box"),
            ("Total de Estudios / Prácticas Realizadas", int(m.get('total_estudios', 0)), "Prácticas efectivas (excluye canceladas)"),
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Indicador Asistencial", "Valor", "Descripción Operativa"],
            filas_kpis,
            alineaciones=['L', 'R', 'L']
        )

        # 2. Distribución y SLA por Categoría de Triage
        fila = cls._agregar_titulo_seccion(ws, fila, "2. Distribución y Cumplimiento de SLA por Categoría de Triage", max_cols=6)
        total_emitidos = int(m.get('total_emitidos', 0) or 0)
        sla_map = {item['clasificacion_triage']: item for item in m.get('sla_triage', [])}
        dist_map = {item['clasificacion_triage']: int(item.get('cantidad', 0)) for item in m.get('distribucion_triage', [])}

        filas_triage = []
        fila_inicio_sla = fila + 1
        for key_triage, label_triage in ETIQUETAS_TRIAGE.items():
            sla_item = sla_map.get(key_triage, {})
            cant_pac = int(sla_item.get('total_pacientes', dist_map.get(key_triage, 0)) or 0)
            pct_triage = round((cant_pac / total_emitidos) * 100, 1) if total_emitidos > 0 else 0.0
            esp_prom = float(sla_item.get('promedio_espera_minutos', 0.0) or 0.0)
            obj_min = OBJETIVOS_SLA_TRIAGE.get(key_triage)
            obj_txt = f"{int(obj_min)} min" if obj_min is not None else "Sin tope fijo"

            if cant_pac == 0 or esp_prom <= 0:
                estado_sla = "Sin datos"
            elif obj_min is None:
                estado_sla = "Informativo"
            elif esp_prom <= obj_min:
                estado_sla = "Cumple SLA"
            else:
                estado_sla = "Excedido"

            filas_triage.append((label_triage, cant_pac, pct_triage, esp_prom, obj_txt, estado_sla))

        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Categoría de Triage", "Pacientes", "% del Total", "Espera Prom. (min)", "Objetivo SLA", "Estado SLA"],
            filas_triage,
            alineaciones=['L', 'R', 'R', 'R', 'C', 'C']
        )
        # Colorear semáforo de Estado SLA
        for idx_r in range(len(filas_triage)):
            celda_estado = ws.cell(row=fila_inicio_sla + idx_r, column=6)
            if celda_estado.value == "Cumple SLA":
                celda_estado.fill = cls._FILL_OK
                celda_estado.font = cls._FONT_OK
            elif celda_estado.value == "Excedido":
                celda_estado.fill = cls._FILL_ALERT
                celda_estado.font = cls._FONT_ALERT

        # 3. Franjas Horarias Hospitalarias
        fila = cls._agregar_titulo_seccion(ws, fila, "3. Flujo Asistencial por Franja Horaria", max_cols=6)
        filas_franjas = [
            (
                f.get('nombre', '—'),
                f.get('rango_display', '—'),
                int(f.get('total_emitidos', 0)),
                int(f.get('total_atendidos', 0)),
                int(f.get('total_cancelados', 0)),
                float(f.get('promedio_espera_minutos', 0.0)),
                float(f.get('promedio_atencion_minutos', 0.0)),
            )
            for f in m.get('franjas_horarias', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Turno / Franja", "Horario", "Emitidos", "Atendidos", "Cancelados", "Espera Prom. (min)", "Atención Prom. (min)"],
            filas_franjas,
            alineaciones=['L', 'C', 'R', 'R', 'R', 'R', 'R']
        )

        # 4. Curva de Demanda Hora por Hora
        fila = cls._agregar_titulo_seccion(ws, fila, "4. Curva de Demanda Hora por Hora", max_cols=6)
        filas_horas = [
            (
                h.get('franja_hora', '—'),
                int(h.get('total_emitidos', 0)),
                int(h.get('total_atendidos', 0)),
                int(h.get('total_criticos', 0)),
                int(h.get('total_cancelados', 0)),
                float(h.get('porcentaje_dia', 0.0)),
            )
            for h in m.get('demanda_por_hora', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Hora de Admisión", "Tickets Emitidos", "Atendidos", "Críticos Atendidos", "Cancelados", "% del Flujo Diario"],
            filas_horas,
            alineaciones=['C', 'R', 'R', 'R', 'R', 'R'],
            mensaje_vacio="Sin ingresos registrados por hora en la jornada."
        )

        # 5. Resumen por Obra Social / Cobertura Médica
        fila = cls._agregar_titulo_seccion(ws, fila, "5. Distribución por Obra Social / Cobertura Médica", max_cols=6)
        filas_os = [
            (
                os_item.get('obra_social', 'Sin especificar'),
                int(os_item.get('total_tickets', 0)),
                int(os_item.get('pacientes_atendidos', 0)),
                int(os_item.get('pacientes_cancelados', 0)),
                int(os_item.get('total_estudios', 0)),
                float(os_item.get('porcentaje_tickets', 0.0)),
            )
            for os_item in m.get('resumen_obras_sociales', [])
        ]
        cls._agregar_tabla(
            ws,
            fila,
            ["Obra Social / Cobertura", "Tickets Emitidos", "Pacientes Atendidos", "Cancelados", "Estudios Realizados", "% de Pacientes"],
            filas_os,
            alineaciones=['L', 'R', 'R', 'R', 'R', 'R'],
            mensaje_vacio="Sin datos de cobertura médica registrados en la jornada."
        )

        cls._autoajustar_columnas(ws)

    @classmethod
    def _construir_hoja_estudios_e_insumos(
        cls,
        ws: Worksheet,
        m: Dict[str, Any],
        fecha_txt: str,
        generado_txt: str
    ) -> None:
        ws.title = "2. Estudios e Insumos"
        fila = cls._agregar_encabezado_hoja(
            ws,
            "Hoja 2 · Volumen por Sección Bioquímica, Logística de Muestras y Detalle de Prácticas",
            fecha_txt,
            generado_txt,
            max_cols=5
        )

        secciones = m.get('resumen_secciones', [])
        # 1. Resumen por sección
        fila = cls._agregar_titulo_seccion(ws, fila, "1. Resumen de Estudios por Sección de Laboratorio", max_cols=5)
        filas_sec = [
            (
                s.get('nombre_seccion', '—'),
                int(s.get('total_estudios', 0)),
                float(s.get('porcentaje', 0.0)),
                len(s.get('practicas', [])),
            )
            for s in secciones
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Sección Bioquímica", "Total Estudios Realizados", "% del Total", "Prácticas Distintas"],
            filas_sec,
            alineaciones=['L', 'R', 'R', 'R'],
            mensaje_vacio="Sin estudios realizados en la jornada."
        )

        # 2. Insumos y Tubos
        fila = cls._agregar_titulo_seccion(ws, fila, "2. Logística de Muestras, Tubos e Insumos Requeridos", max_cols=5)
        filas_insumos = [
            (
                str(ins.get('tipo_muestra', '—')).upper(),
                ins.get('descripcion_tubo', '—'),
                int(ins.get('total_tubos', 0)),
                float(ins.get('porcentaje', 0.0)),
            )
            for ins in m.get('insumos_tubos', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Código Muestra", "Descripción del Tubo / Contenedor", "Unidades Utilizadas", "% del Total"],
            filas_insumos,
            alineaciones=['C', 'L', 'R', 'R'],
            mensaje_vacio="Sin insumos registrados en la jornada."
        )

        # 3. Detalle de todas las prácticas
        fila = cls._agregar_titulo_seccion(ws, fila, "3. Catálogo Nominal de Prácticas Realizadas en el Día", max_cols=5)
        filas_practicas = []
        for s in secciones:
            nombre_sec = s.get('nombre_seccion', '—')
            for pr in s.get('practicas', []):
                filas_practicas.append((
                    nombre_sec,
                    pr.get('codigo') or '—',
                    pr.get('nombre', '—'),
                    str(pr.get('tipo_muestra', '—')).capitalize(),
                    int(pr.get('total', 0)),
                ))

        cls._agregar_tabla(
            ws,
            fila,
            ["Sección Bioquímica", "Código Práctica", "Nombre de la Práctica", "Tipo de Muestra", "Cantidad Realizada"],
            filas_practicas,
            alineaciones=['L', 'C', 'L', 'C', 'R'],
            mensaje_vacio="No se registraron prácticas bioquímicas en la jornada."
        )

        cls._autoajustar_columnas(ws)

    @classmethod
    def _construir_hoja_personal_y_boxes(
        cls,
        ws: Worksheet,
        m: Dict[str, Any],
        fecha_txt: str,
        generado_txt: str
    ) -> None:
        ws.title = "3. Personal y Boxes"
        fila = cls._agregar_encabezado_hoja(
            ws,
            "Hoja 3 · Productividad de Admisión, Técnicos de Extracción y Uso de Boxes Físicos",
            fecha_txt,
            generado_txt,
            max_cols=7
        )

        # 1. Operadores de Admisión
        fila = cls._agregar_titulo_seccion(ws, fila, "1. Rendimiento de Operadores de Admisión", max_cols=7)
        filas_adm = [
            (
                adm.get('nombre_completo', '—'),
                adm.get('personal_admision__dni') or '—',
                adm.get('personal_admision__matricula') or '—',
                adm.get('personal_admision__email') or '—',
                int(adm.get('pacientes_recepcionados', 0)),
            )
            for adm in m.get('rendimiento_admision', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Operador de Admisión", "DNI", "Matrícula", "Correo Institucional", "Tickets Recepcionados"],
            filas_adm,
            alineaciones=['L', 'C', 'C', 'L', 'R'],
            mensaje_vacio="Sin admisiones registradas en la jornada."
        )

        # 2. Técnicos de Box
        fila = cls._agregar_titulo_seccion(ws, fila, "2. Rendimiento de Técnicos de Box de Extracción", max_cols=7)
        filas_box_personal = [
            (
                bx.get('nombre_completo', '—'),
                bx.get('dni') or '—',
                bx.get('matricula') or '—',
                bx.get('email') or '—',
                int(bx.get('pacientes_atendidos', 0)),
                int(bx.get('pacientes_ausentes', 0)),
                int(bx.get('total_convocados', 0)),
                float(bx.get('promedio_atencion_minutos', 0.0)),
            )
            for bx in m.get('rendimiento_box', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            [
                "Técnico de Box",
                "DNI",
                "Matrícula",
                "Correo Institucional",
                "Pacientes Atendidos",
                "Pacientes Ausentes",
                "Total Convocados",
                "Atención Prom. (min)",
            ],
            filas_box_personal,
            alineaciones=['L', 'C', 'C', 'L', 'R', 'R', 'R', 'R'],
            mensaje_vacio="Sin atenciones en box registradas en la jornada."
        )

        # 3. Uso y Rendimiento por Box Físico
        fila = cls._agregar_titulo_seccion(ws, fila, "3. Ocupación y Rendimiento por Box Físico de Extracción", max_cols=7)
        filas_boxes_fisicos = [
            (
                b.get('nombre_box', f"Box {b.get('numero_box', '—')}"),
                "Sí (Prioritario)" if b.get('es_discapacidad') else "No (Estándar)",
                int(b.get('total_convocados', 0)),
                int(b.get('pacientes_atendidos', 0)),
                int(b.get('pacientes_ausentes', 0)),
                int(b.get('pacientes_discapacidad_atendidos', 0)),
                float(b.get('promedio_atencion_minutos', 0.0)),
            )
            for b in m.get('rendimiento_boxes_fisicos', [])
        ]
        cls._agregar_tabla(
            ws,
            fila,
            [
                "Box Físico",
                "Accesibilidad / Discapacidad",
                "Total Convocados",
                "Pacientes Atendidos",
                "Pacientes Ausentes",
                "Atendidos Movilidad Reducida",
                "Atención Prom. (min)",
            ],
            filas_boxes_fisicos,
            alineaciones=['L', 'C', 'R', 'R', 'R', 'R', 'R'],
            mensaje_vacio="Sin actividad en boxes físicos durante la jornada."
        )

        cls._autoajustar_columnas(ws)

    @classmethod
    def _construir_hoja_horas_extras(
        cls,
        ws: Worksheet,
        m: Dict[str, Any],
        fecha_txt: str,
        generado_txt: str
    ) -> None:
        ws.title = "4. Horas Extras"
        fila = cls._agregar_encabezado_hoja(
            ws,
            "Hoja 4 · Control de Horas Extras, Extensiones de Turno y Habilitaciones Excepcionales",
            fecha_txt,
            generado_txt,
            max_cols=10
        )

        horas_extras = m.get('horas_extras_dia', [])
        total_hab = len(horas_extras)
        total_hs = round(sum(float(item.get('cantidad_horas', 0.0) or 0.0) for item in horas_extras), 2)

        # 1. Resumen de horas extras del día
        fila = cls._agregar_titulo_seccion(ws, fila, "1. Resumen de Horas Extras Autorizadas en la Jornada", max_cols=10)
        filas_resumen_he = [
            ("Total de Habilitaciones Horarias Activas", total_hab),
            ("Total de Horas Extras Autorizadas (hs)", total_hs),
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Indicador de Recursos Humanos", "Total"],
            filas_resumen_he,
            alineaciones=['L', 'R']
        )

        # 2. Nómina detallada de horas extras
        fila = cls._agregar_titulo_seccion(ws, fila, "2. Nómina Detallada de Horas Extras y Excepciones de Turno", max_cols=10)
        filas_detalle_he = [
            (
                he.get('personal_nombre', '—'),
                he.get('personal_dni', '—'),
                he.get('personal_matricula', '—'),
                he.get('personal_rol', '—'),
                he.get('horario_habitual', '—'),
                he.get('franja_autorizada', '—'),
                float(he.get('cantidad_horas', 0.0)),
                he.get('tipo_extension', '—'),
                he.get('motivo', '—'),
                he.get('aprobado_por', '—'),
            )
            for he in horas_extras
        ]
        cls._agregar_tabla(
            ws,
            fila,
            [
                "Personal Autorizado",
                "DNI",
                "Matrícula",
                "Rol Operativo",
                "Turno Habitual",
                "Franja Extra Autorizada",
                "Horas Computadas",
                "Tipo de Extensión",
                "Motivo / Justificación",
                "Autorizado Por",
            ],
            filas_detalle_he,
            alineaciones=['L', 'C', 'C', 'C', 'C', 'C', 'R', 'L', 'L', 'L'],
            mensaje_vacio="No se registraron horas extras ni habilitaciones excepcionales para esta fecha."
        )

        cls._autoajustar_columnas(ws)

    @classmethod
    def _construir_hoja_auditoria(
        cls,
        ws: Worksheet,
        m: Dict[str, Any],
        fecha_txt: str,
        generado_txt: str
    ) -> None:
        ws.title = "5. Auditoría"
        fila = cls._agregar_encabezado_hoja(
            ws,
            "Hoja 5 · Auditoría de Modificaciones en Box, Cancelaciones y Justificaciones de Triage",
            fecha_txt,
            generado_txt,
            max_cols=9
        )

        auditoria = m.get('auditoria_box', {}) or {}

        # 1. KPIs de Rectificación en Box
        fila = cls._agregar_titulo_seccion(ws, fila, "1. Indicadores de Calidad y Rectificación de Estudios en Box", max_cols=9)
        filas_kpi_audit = [
            ("Estudios Cargados Inicialmente en Admisión", int(auditoria.get('total_cargados_admision', 0))),
            ("Estudios Agregados en Box de Extracción (+)", int(auditoria.get('total_agregados_box', 0))),
            ("Estudios Cancelados en Box de Extracción (-)", int(auditoria.get('total_cancelados_box', 0))),
            ("Tasa de Rectificación Operativa (%)", float(auditoria.get('tasa_rectificacion_porcentaje', 0.0))),
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Métrica de Auditoría", "Valor"],
            filas_kpi_audit,
            alineaciones=['L', 'R']
        )

        # 2. Motivos frecuentes de cancelación de estudios
        fila = cls._agregar_titulo_seccion(ws, fila, "2. Consolidado de Motivos de Cancelación de Estudios", max_cols=9)
        filas_motivos = [
            (
                mot.get('motivo_modificacion', 'Sin especificar'),
                int(mot.get('cantidad', 0)),
            )
            for mot in auditoria.get('motivos_cancelacion', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            ["Motivo de Cancelación Registrado", "Cantidad de Estudios"],
            filas_motivos,
            alineaciones=['L', 'R'],
            mensaje_vacio="Sin estudios cancelados con motivo registrado en la jornada."
        )

        # 3. Trazabilidad nominal de modificaciones de estudios en Box
        fila = cls._agregar_titulo_seccion(
            ws,
            fila,
            "3. Trazabilidad Nominal de Estudios Agregados o Cancelados en Box",
            max_cols=9
        )
        filas_mod_estudios = [
            (
                mod.get('num_totem', '—'),
                mod.get('codigo_practica', '—'),
                mod.get('nombre_estudio', '—'),
                mod.get('seccion', '—'),
                mod.get('accion', '—'),
                mod.get('box', '—'),
                mod.get('tecnico', '—'),
                mod.get('hora', '—'),
                mod.get('motivo', '—'),
            )
            for mod in auditoria.get('detalle_modificaciones_estudios', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            [
                "N° Ticket",
                "Código",
                "Estudio / Práctica",
                "Sección",
                "Acción en Box",
                "Box",
                "Técnico Responsable",
                "Hora",
                "Motivo / Observaciones",
            ],
            filas_mod_estudios,
            alineaciones=['C', 'C', 'L', 'L', 'C', 'C', 'L', 'C', 'L'],
            mensaje_vacio="No se agregaron ni cancelaron estudios en box durante la jornada."
        )

        # 4. Cierres de Box Cancelados / Ausentes o con Observaciones
        fila = cls._agregar_titulo_seccion(
            ws,
            fila,
            "4. Auditoría de Atenciones Canceladas, Ausencias u Observaciones de Cierre en Box",
            max_cols=9
        )
        filas_cierres = [
            (
                c.get('num_totem', '—'),
                c.get('box', '—'),
                c.get('tecnico', '—'),
                c.get('hora_llamado', '—'),
                c.get('hora_cierre', '—'),
                c.get('motivo_cierre', '—'),
                c.get('observaciones_cierre', '—'),
            )
            for c in auditoria.get('detalle_cierres_cancelados_box', [])
        ]
        fila = cls._agregar_tabla(
            ws,
            fila,
            [
                "N° Ticket",
                "Box",
                "Técnico a Cargo",
                "Hora Llamado",
                "Hora Cierre",
                "Motivo de Cierre",
                "Observaciones Registradas",
            ],
            filas_cierres,
            alineaciones=['C', 'C', 'L', 'C', 'C', 'C', 'L'],
            mensaje_vacio="Sin atenciones canceladas ni observaciones especiales de cierre en box."
        )

        # 5. Auditoría de Triage "Otro"
        fila = cls._agregar_titulo_seccion(
            ws,
            fila,
            "5. Auditoría de Admisiones con Clasificación de Triage «Otro» (Justificación Médica Obligatoria)",
            max_cols=9
        )
        filas_triage_otro = [
            (
                to.get('num_totem', '—'),
                to.get('hora_admision', '—'),
                to.get('dni_paciente', '—'),
                to.get('obra_social', '—'),
                to.get('operador_admision', '—'),
                to.get('justificacion_otro', '—'),
                to.get('estado', '—'),
            )
            for to in auditoria.get('detalle_triage_otro', [])
        ]
        cls._agregar_tabla(
            ws,
            fila,
            [
                "N° Ticket",
                "Hora Admisión",
                "DNI Paciente",
                "Obra Social",
                "Operador de Admisión",
                "Justificación Médica Registrada",
                "Estado del Ticket",
            ],
            filas_triage_otro,
            alineaciones=['C', 'C', 'C', 'L', 'L', 'L', 'C'],
            mensaje_vacio="No se emitieron tickets bajo la categoría «Otro» en la jornada."
        )

        cls._autoajustar_columnas(ws)

    @classmethod
    def _construir_hoja_detalle_tickets(
        cls,
        ws: Worksheet,
        m: Dict[str, Any],
        fecha_txt: str,
        generado_txt: str
    ) -> None:
        ws.title = "6. Detalle de Tickets"
        fila = cls._agregar_encabezado_hoja(
            ws,
            "Hoja 6 · Listado Nominal Completo de Tickets y Tiempos de Atención de la Jornada",
            fecha_txt,
            generado_txt,
            max_cols=16
        )

        fila = cls._agregar_titulo_seccion(
            ws,
            fila,
            "Sábana Operativa de Tickets Admitidos en la Jornada",
            max_cols=16
        )

        filas_tickets = [
            (
                t.get('num_totem', '—'),
                t.get('hora_admision', '—'),
                t.get('dni_paciente', '—'),
                t.get('obra_social', '—'),
                t.get('clasificacion_triage_display', '—'),
                t.get('justificacion_otro', '—'),
                t.get('operador_admision', '—'),
                t.get('estado', '—'),
                t.get('box_atencion', '—'),
                t.get('tecnico_box', '—'),
                t.get('hora_llamado', '—'),
                t.get('hora_finalizacion', '—'),
                t.get('espera_minutos') if t.get('espera_minutos') is not None else '—',
                t.get('atencion_minutos') if t.get('atencion_minutos') is not None else '—',
                int(t.get('cantidad_estudios', 0)),
                t.get('estudios_detalle', '—'),
            )
            for t in m.get('detalle_tickets', [])
        ]

        cls._agregar_tabla(
            ws,
            fila,
            [
                "N° Ticket",
                "Hora Admisión",
                "DNI Paciente",
                "Obra Social",
                "Categoría Triage",
                "Justificación (Triage Otro)",
                "Operador Admisión",
                "Estado Ticket",
                "Box Atención",
                "Técnico Box",
                "Hora Llamado",
                "Hora Cierre",
                "Espera (min)",
                "Atención (min)",
                "Cant. Estudios",
                "Estudios Realizados",
            ],
            filas_tickets,
            alineaciones=['C', 'C', 'C', 'L', 'L', 'L', 'L', 'C', 'C', 'L', 'C', 'C', 'R', 'R', 'R', 'L'],
            mensaje_vacio="No se emitieron tickets asistenciales durante la jornada seleccionada."
        )

        cls._autoajustar_columnas(ws, min_width=12, max_width=52)

    @classmethod
    def generar(cls, metricas: Dict[str, Any]) -> bytes:
        """
        Compila el libro oficial Excel (.xlsx) con las 6 hojas temáticas y retorna sus bytes en memoria.
        """
        wb = Workbook()
        fecha_txt = _fecha_legible(metricas.get('fecha', ''))
        generado_txt = timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M')

        ws1 = wb.active
        cls._construir_hoja_resumen_y_flujo(ws1, metricas, fecha_txt, generado_txt)

        ws2 = wb.create_sheet()
        cls._construir_hoja_estudios_e_insumos(ws2, metricas, fecha_txt, generado_txt)

        ws3 = wb.create_sheet()
        cls._construir_hoja_personal_y_boxes(ws3, metricas, fecha_txt, generado_txt)

        ws4 = wb.create_sheet()
        cls._construir_hoja_horas_extras(ws4, metricas, fecha_txt, generado_txt)

        ws5 = wb.create_sheet()
        cls._construir_hoja_auditoria(ws5, metricas, fecha_txt, generado_txt)

        ws6 = wb.create_sheet()
        cls._construir_hoja_detalle_tickets(ws6, metricas, fecha_txt, generado_txt)

        buffer = io.BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        buffer.close()
        return excel_bytes


__all__ = ['ReportExcelGenerator']

